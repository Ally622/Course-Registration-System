"""
Admin export routes — PDF and Excel report generation.
"""

import io
from datetime import datetime

from flask import request, jsonify, send_file

from admin import admin_bp
from db import get_connection
from auth.decorators import admin_required

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle,
    Paragraph, Spacer
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# =============================================================
# PDF EXPORT — Student List
# =============================================================
@admin_bp.route('/admin/export/students/pdf', methods=['GET'])
@admin_required
def export_students_pdf():
    """Export all students as a PDF report."""
    connection = get_connection()
    cursor = connection.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT
                s.registration_number,
                s.student_name,
                u.email,
                s.phone,
                d.department_name,
                s.year_of_study
            FROM students s
            INNER JOIN users u ON s.user_id = u.user_id
            LEFT JOIN departments d ON s.department_id = d.department_id
            ORDER BY s.student_name
        """)
        students = cursor.fetchall()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), topMargin=0.5 * inch)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ReportTitle', parent=styles['Title'],
            fontSize=16, textColor=colors.HexColor('#C2185B'),
            alignment=TA_CENTER
        )

        elements.append(Paragraph("Student Report", title_style))
        elements.append(Paragraph(
            f"Generated on {datetime.now().strftime('%B %d, %Y')}",
            ParagraphStyle('Sub', parent=styles['Normal'], alignment=TA_CENTER)
        ))
        elements.append(Spacer(1, 20))

        data = [["#", "Reg. Number", "Name", "Email", "Phone", "Department", "Year"]]
        for i, s in enumerate(students, 1):
            data.append([
                str(i),
                s['registration_number'],
                s['student_name'],
                s['email'],
                s.get('phone', ''),
                s.get('department_name', 'N/A'),
                str(s.get('year_of_study', ''))
            ])

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#C2185B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF0F5')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)

        doc.build(elements)
        buffer.seek(0)

        return send_file(
            buffer, as_attachment=True,
            download_name=f"students_report_{datetime.now().strftime('%Y%m%d')}.pdf",
            mimetype='application/pdf'
        )

    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

    finally:
        cursor.close()
        connection.close()


# =============================================================
# EXCEL EXPORT — Student List
# =============================================================
@admin_bp.route('/admin/export/students/excel', methods=['GET'])
@admin_required
def export_students_excel():
    """Export all students as an Excel file."""
    connection = get_connection()
    cursor = connection.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT
                s.registration_number,
                s.student_name,
                u.email,
                s.phone,
                d.department_name,
                s.year_of_study
            FROM students s
            INNER JOIN users u ON s.user_id = u.user_id
            LEFT JOIN departments d ON s.department_id = d.department_id
            ORDER BY s.student_name
        """)
        students = cursor.fetchall()

        wb = Workbook()
        ws = wb.active
        ws.title = "Students"

        # Styles
        header_fill = PatternFill(start_color="C2185B", end_color="C2185B", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        headers = ["#", "Reg. Number", "Name", "Email", "Phone", "Department", "Year"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        for i, s in enumerate(students, 1):
            row = i + 1
            values = [
                i,
                s['registration_number'],
                s['student_name'],
                s['email'],
                s.get('phone', ''),
                s.get('department_name', 'N/A'),
                s.get('year_of_study', '')
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.border = thin_border

        # Auto-size columns
        for col in range(1, len(headers) + 1):
            max_len = max(
                len(str(ws.cell(row=r, column=col).value or ''))
                for r in range(1, ws.max_row + 1)
            )
            ws.column_dimensions[chr(64 + col)].width = max_len + 3

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return send_file(
            buffer, as_attachment=True,
            download_name=f"students_report_{datetime.now().strftime('%Y%m%d')}.xlsx",
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

    finally:
        cursor.close()
        connection.close()


# =============================================================
# PDF EXPORT — Registrations Report
# =============================================================
@admin_bp.route('/admin/export/registrations/pdf', methods=['GET'])
@admin_required
def export_registrations_pdf():
    """Export registration report as PDF."""
    connection = get_connection()
    cursor = connection.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT
                s.registration_number,
                s.student_name,
                c.course_code,
                c.course_name,
                r.status,
                r.registration_date
            FROM registrations r
            INNER JOIN students s ON r.student_id = s.student_id
            INNER JOIN courses c ON r.course_id = c.course_id
            INNER JOIN semesters sem ON r.semester_id = sem.semester_id
            WHERE sem.is_active = 1
            ORDER BY s.student_name, c.course_code
        """)
        registrations = cursor.fetchall()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), topMargin=0.5 * inch)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ReportTitle', parent=styles['Title'],
            fontSize=16, textColor=colors.HexColor('#C2185B'),
            alignment=TA_CENTER
        )

        elements.append(Paragraph("Registration Report", title_style))
        elements.append(Spacer(1, 20))

        data = [["#", "Reg. Number", "Student", "Course Code", "Course Name", "Status", "Date"]]
        for i, r in enumerate(registrations, 1):
            reg_date = r['registration_date'].strftime('%Y-%m-%d') if r.get('registration_date') else ''
            data.append([
                str(i),
                r['registration_number'],
                r['student_name'],
                r['course_code'],
                r['course_name'],
                r['status'].capitalize(),
                reg_date
            ])

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#C2185B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF0F5')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)

        doc.build(elements)
        buffer.seek(0)

        return send_file(
            buffer, as_attachment=True,
            download_name=f"registrations_report_{datetime.now().strftime('%Y%m%d')}.pdf",
            mimetype='application/pdf'
        )

    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

    finally:
        cursor.close()
        connection.close()


# =============================================================
# EXCEL EXPORT — Registrations Report
# =============================================================
@admin_bp.route('/admin/export/registrations/excel', methods=['GET'])
@admin_required
def export_registrations_excel():
    """Export registration report as Excel."""
    connection = get_connection()
    cursor = connection.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT
                s.registration_number,
                s.student_name,
                c.course_code,
                c.course_name,
                r.status,
                r.registration_date
            FROM registrations r
            INNER JOIN students s ON r.student_id = s.student_id
            INNER JOIN courses c ON r.course_id = c.course_id
            INNER JOIN semesters sem ON r.semester_id = sem.semester_id
            WHERE sem.is_active = 1
            ORDER BY s.student_name, c.course_code
        """)
        registrations = cursor.fetchall()

        wb = Workbook()
        ws = wb.active
        ws.title = "Registrations"

        header_fill = PatternFill(start_color="C2185B", end_color="C2185B", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        headers = ["#", "Reg. Number", "Student", "Course Code", "Course Name", "Status", "Date"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        for i, r in enumerate(registrations, 1):
            row = i + 1
            reg_date = r['registration_date'].strftime('%Y-%m-%d') if r.get('registration_date') else ''
            values = [
                i,
                r['registration_number'],
                r['student_name'],
                r['course_code'],
                r['course_name'],
                r['status'].capitalize(),
                reg_date
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.border = thin_border

        for col in range(1, len(headers) + 1):
            max_len = max(
                len(str(ws.cell(row=r, column=col).value or ''))
                for r in range(1, ws.max_row + 1)
            )
            adjusted = min(max_len + 3, 40)
            ws.column_dimensions[chr(64 + col)].width = adjusted

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return send_file(
            buffer, as_attachment=True,
            download_name=f"registrations_report_{datetime.now().strftime('%Y%m%d')}.xlsx",
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

    finally:
        cursor.close()
        connection.close()
